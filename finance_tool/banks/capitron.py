import re
from datetime import date, datetime
from decimal import Decimal
from openpyxl import load_workbook
from finance_tool.models import ParsedStatement, ParsedTransaction, StatementError

def money(value):
    return Decimal(str(value or 0).replace(',','').replace(' MNT','').strip())

class CapitronAdapter:
    code='capitron'
    def detect(self,path):
        try:
            w=load_workbook(path,read_only=True,data_only=True)
            s=w.active
            good=s['A5'].value=='Дансны дугаар' and s['A10'].value=='Огноо' and s['G10'].value=='Гүйлгээний утга'
            w.close()
            return good
        except Exception:
            return False
    def parse(self,path):
        if not self.detect(path): raise StatementError('Capitron хуулгын загвар танигдсангүй (A5/A10/G10).')
        w=load_workbook(path,read_only=True,data_only=True);s=w.active
        try:
            period=str(s['G5'].value).split(' - ')
            st=ParsedStatement(bank_code='capitron',account_no=str(s['B5'].value),holder_name=str(s['B6'].value),period_from=date.fromisoformat(period[0]),period_to=date.fromisoformat(period[1]),opening_balance=money(s['B7'].value),closing_balance=money(s['B8'].value))
            tx=[]
            for rowno,row in enumerate(s.iter_rows(min_row=11,values_only=True),11):
                a,b,c,d,e,f,g,h=tuple(row[:8])
                if a=='Нийт төлбөр':
                    st.file_debit_total=money(b);st.file_credit_total=money(c);break
                if not a: continue
                try:
                    description=str(g or '').strip()
                    acct=str(e or '').strip() or None
                    name=None
                    if acct and not (acct.startswith('MN') or acct.isdigit()):name=acct
                    tx.append(ParsedTransaction(date.fromisoformat(str(a)[:10]),datetime.fromisoformat(str(h)),money(b),money(c),money(f),description,acct,name))
                except Exception as exc:raise StatementError(f'Хуулгын {rowno}-р мөрийн өгөгдөл буруу: {exc}') from exc
            tx.reverse();st.transactions=tx
            if not tx:raise StatementError('Хуулгад гүйлгээ алга.')
            if st.file_debit_total is not None and sum(t.debit for t in tx)!=st.file_debit_total:raise StatementError('Зарлагын нийлбэр банкны дүнтэй таарахгүй.')
            if st.file_credit_total is not None and sum(t.credit for t in tx)!=st.file_credit_total:raise StatementError('Орлогын нийлбэр банкны дүнтэй таарахгүй.')
            balance=st.opening_balance
            for t in tx:
                balance+=t.credit-t.debit
                if balance!=t.balance_after:
                    # Original sheet is newest first, locate matching row for useful error.
                    original=next((i for i,r in enumerate(s.iter_rows(min_row=11,values_only=True),11) if str(r[7])==t.posted_at.isoformat(sep=' ') and money(r[5])==t.balance_after),None)
                    raise StatementError(f'Үлдэгдэл зөрсөн (Excel мөр {original or "тодорхойгүй"}, {t.posted_at}).')
            if balance!=st.closing_balance:raise StatementError('Эцсийн үлдэгдэл таарахгүй.')
            return st
        finally:w.close()
