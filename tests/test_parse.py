"""Checks the Capitron parser against the anonymised sample and its expected numbers.

Run from the project folder:  pytest tests/test_parse.py -q
"""
import json
import shutil
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from finance_tool.banks.capitron import CapitronAdapter
from finance_tool.models import StatementError

ROOT = Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "fixtures" / "capitron_sample_anonymised.xlsx"
EXPECTED = json.loads((ROOT / "fixtures" / "expected_capitron_sample.json").read_text(encoding="utf-8"))


def test_detects_capitron_file():
    assert CapitronAdapter().detect(str(FIXTURE)) is True


def test_header_fields():
    st = CapitronAdapter().parse(str(FIXTURE))
    assert st.bank_code == "capitron"
    assert st.account_no == EXPECTED["header"]["account_no"]
    assert st.holder_name == EXPECTED["header"]["holder_name"]
    assert st.period_from == date(2026, 7, 14)
    assert st.period_to == date(2026, 10, 14)
    assert st.opening_balance == Decimal("22742")
    assert st.closing_balance == Decimal("2092")


def test_row_count_and_totals():
    st = CapitronAdapter().parse(str(FIXTURE))
    assert len(st.transactions) == EXPECTED["rows_in_file"] == 107
    assert sum(t.debit for t in st.transactions) == Decimal(EXPECTED["debit_total"])
    assert sum(t.credit for t in st.transactions) == Decimal(EXPECTED["credit_total"])
    assert st.file_debit_total == Decimal(EXPECTED["file_totals_row"]["debit"])
    assert st.file_credit_total == Decimal(EXPECTED["file_totals_row"]["credit"])


def test_rows_are_oldest_first_and_fields_are_read():
    st = CapitronAdapter().parse(str(FIXTURE))
    oldest, newest = st.transactions[0], st.transactions[-1]
    assert oldest.posted_at == datetime(2026, 8, 8, 19, 10, 27)
    assert oldest.credit == Decimal("1005000")
    assert oldest.balance_after == Decimal("1027742")
    assert newest.posted_at == datetime(2026, 10, 4, 1, 51, 35)
    assert newest.debit == Decimal("50")
    assert newest.txn_date == date(2026, 10, 4)


def test_amounts_are_decimal_not_float():
    st = CapitronAdapter().parse(str(FIXTURE))
    for t in st.transactions:
        assert isinstance(t.debit, Decimal) and isinstance(t.credit, Decimal)
        assert isinstance(t.balance_after, Decimal)


def test_description_text_is_kept_exactly_after_trim():
    st = CapitronAdapter().parse(str(FIXTURE))
    texts = [t.description for t in st.transactions]
    assert any(t.startswith("POS:") for t in texts)
    assert any(t.startswith("qpay ") for t in texts)
    assert all(t == t.strip() for t in texts)


def test_broken_balance_is_rejected_with_row_number(tmp_path):
    bad = tmp_path / "broken.xlsx"
    shutil.copy(FIXTURE, bad)
    wb = openpyxl.load_workbook(bad)
    ws = wb.worksheets[0]
    ws["F50"].value = int(ws["F50"].value) + 1  # balance column, sheet row 50
    wb.save(bad)
    with pytest.raises(StatementError) as err:
        CapitronAdapter().parse(str(bad))
    assert "49" in str(err.value) or "50" in str(err.value)


def test_non_statement_file_is_not_detected(tmp_path):
    other = tmp_path / "other.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "hello"
    wb.save(other)
    assert CapitronAdapter().detect(str(other)) is False
